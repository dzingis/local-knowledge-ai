from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from backend.parsers.parser_models import ParsedDocument, ParsedObjectType, ParserKind, ParserReference, ParserSummary


@dataclass(frozen=True, slots=True)
class WordBlock:
    index: int
    text: str
    style: str | None = None


@dataclass(frozen=True, slots=True)
class WordParseResult:
    source_path: Path
    title: str
    blocks: tuple[WordBlock, ...]
    text: str
    page_count: int | None = None
    metadata: dict[str, str] | None = None


class WordParser:
    """Parse Word documents into a normalized, search-ready representation."""

    def parse_file(self, path: str | Path) -> ParsedDocument:
        source_path = Path(path)
        text, title, metadata, page_count = self._extract_word_text(source_path)
        lines = text.splitlines()
        references = self._extract_references(lines)
        return ParsedDocument(
            source_path=source_path,
            parser_kind=ParserKind.WORD,
            object_type=ParsedObjectType.DOCUMENT,
            title=title,
            text=text,
            metadata={
                "source_format": "docx",
                "word_title": title,
                **metadata,
            },
            references=references,
            checksum=self._checksum_text(text),
            language="unknown",
            page_count=page_count,
            line_count=len(lines),
            character_count=len(text),
        )

    def parse_many(self, paths: Iterable[str | Path]) -> tuple[ParsedDocument, ...]:
        return tuple(self.parse_file(path) for path in paths)

    def summarize(self, documents: Iterable[ParsedDocument]) -> ParserSummary:
        docs = tuple(documents)
        if not docs:
            return ParserSummary(
                source_path=Path("<empty>"),
                parser_kind=ParserKind.WORD,
                object_type=ParsedObjectType.DOCUMENT,
                title="",
                line_count=0,
                character_count=0,
                reference_count=0,
            )
        first = docs[0]
        return ParserSummary(
            source_path=first.source_path,
            parser_kind=ParserKind.WORD,
            object_type=ParsedObjectType.DOCUMENT,
            title=first.title,
            line_count=sum(doc.line_count or 0 for doc in docs),
            character_count=sum(doc.character_count for doc in docs),
            reference_count=sum(len(doc.references) for doc in docs),
        )

    def _extract_word_text(self, path: Path) -> tuple[str, str, dict[str, str], int | None]:
        try:
            from docx import Document  # type: ignore
        except Exception:
            return self._fallback_extract(path)

        try:
            document = Document(str(path))
        except Exception:
            return self._fallback_extract(path)

        paragraphs: list[str] = []
        title = path.stem
        metadata: dict[str, str] = {}
        page_count = None

        core = getattr(document, "core_properties", None)
        if core is not None:
            title = (getattr(core, "title", None) or title).strip() or title
            metadata.update(
                {
                    "author": getattr(core, "author", "") or "",
                    "category": getattr(core, "category", "") or "",
                    "comments": getattr(core, "comments", "") or "",
                    "subject": getattr(core, "subject", "") or "",
                    "keywords": getattr(core, "keywords", "") or "",
                }
            )

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                paragraphs.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))

        text = "\n".join(paragraphs).strip()
        return text, title, metadata, page_count

    def _fallback_extract(self, path: Path) -> tuple[str, str, dict[str, str], int | None]:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return text, path.stem, {"fallback": "plain-text"}, None

    def _extract_references(self, lines: list[str]) -> tuple[ParserReference, ...]:
        references: list[ParserReference] = []
        seen: set[str] = set()
        for index, line in enumerate(lines, start=1):
            for candidate in self._find_markers(line):
                if candidate in seen:
                    continue
                seen.add(candidate)
                references.append(ParserReference(name=candidate, kind="marker", line=index))
        return tuple(references)

    def _find_markers(self, text: str) -> list[str]:
        markers: list[str] = []
        for token in text.replace("/", " ").replace("-", " ").split():
            if token.startswith("[") and token.endswith("]"):
                token = token[1:-1]
            if token and token[0].isupper() and len(token) > 2:
                markers.append(token)
        return markers

    def _checksum_text(self, text: str) -> str:
        import hashlib

        digest = hashlib.sha256()
        digest.update(text.encode("utf-8", errors="ignore"))
        return digest.hexdigest()
